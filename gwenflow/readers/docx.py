import io
from dataclasses import dataclass
from pathlib import Path
from typing import ClassVar, List, Union

from gwenflow.logger import logger
from gwenflow.readers.base import Reader
from gwenflow.types import Document


@dataclass
class DocxReader(Reader):
    trans: ClassVar[dict[int, int | None]] = {
        0x00A0: 0x20,
        0x202F: 0x20,
        0x2007: 0x20,
        0x200B: None,
        0x200C: None,
        0x200D: None,
        0xFEFF: None,
    }

    def get_text(self, file_obj) -> str:
        try:
            import docx
        except ImportError as e:
            raise ImportError("python-docx is not installed. Please install it with `pip install python-docx`") from e
        doc = docx.Document(file_obj)
        return "\n".join((p.text.translate(self.trans) if p.text else "") for p in doc.paragraphs)

    def get_row_cells(self, row):
        from docx.table import _Cell

        try:
            return list(row.cells)
        except (ValueError, IndexError):
            return [_Cell(tc, row.table) for tc in row._tr.tc_lst]

    def get_tables(self, file_obj):
        try:
            import docx
        except ImportError as e:
            raise ImportError("python-docx is not installed. Please install it with `uv add python-docx`") from e
        doc = docx.Document(file_obj)
        tables = []
        for t in doc.tables:
            rows = []
            for r in t.rows:
                cells = []
                for c in self.get_row_cells(r):
                    txt = "\n".join(p.text for p in c.paragraphs) if c.paragraphs else ""
                    txt = txt.translate(self.trans) if txt else ""
                    cells.append(txt)
                rows.append(cells)
            tables.append(rows)
        return tables

    def read(self, file: Union[Path, io.BytesIO]) -> List[Document]:
        try:
            filename = self.get_file_name(file)
            content = self.get_file_content(file)
            data = content.getvalue() if isinstance(content, io.BytesIO) else content
            text = self.get_text(io.BytesIO(data))
            try:
                tables = self.get_tables(io.BytesIO(data))
            except Exception as e:
                logger.warning(f"Could not extract tables from {filename}: {e}")
                tables = []
            table_text = self.format_tables(tables)
            if table_text:
                text = f"{text.strip()}\n\n{table_text}" if text.strip() else table_text

            return [
                Document(
                    id=self.key(filename),
                    content=text,
                    metadata={"filename": filename, "tables": tables},
                )
            ]
        except Exception as e:
            logger.exception(f"Error reading file: {e}")
            return []
